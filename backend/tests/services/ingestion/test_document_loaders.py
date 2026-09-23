from pathlib import Path

import pytest
from docx import Document

from app.services.ingestion.loaders.docx_loader import DOCXLoader
from app.services.ingestion.loaders.pdf_loader import PDFLoader
from app.services.ingestion.loaders.csv_loader import CSVLoader
from app.services.ingestion.document_loader import DocumentLoader

def test_docx_loader_extracts_paragraphs(tmp_path: Path):
    file_path = tmp_path / "policy.docx"

    document = Document()
    document.add_paragraph("Employee Leave Policy")
    document.add_paragraph("Employees must submit leave requests.")
    document.save(file_path)

    loader = DOCXLoader()

    result = loader.load(str(file_path))

    assert "Employee Leave Policy" in result
    assert "Employees must submit leave requests." in result


def test_docx_loader_rejects_non_docx(tmp_path: Path):
    file_path = tmp_path / "policy.txt"
    file_path.write_text("test", encoding="utf-8")

    loader = DOCXLoader()

    with pytest.raises(ValueError, match="Expected a DOCX file"):
        loader.load(str(file_path))


def test_docx_loader_missing_file(tmp_path: Path):
    file_path = tmp_path / "missing.docx"

    loader = DOCXLoader()

    with pytest.raises(FileNotFoundError, match="File not found"):
        loader.load(str(file_path))


def test_pdf_loader_rejects_non_pdf(tmp_path: Path):
    file_path = tmp_path / "policy.txt"
    file_path.write_text("test", encoding="utf-8")

    loader = PDFLoader()

    with pytest.raises(ValueError, match="Expected a PDF file"):
        loader.load(str(file_path))


def test_pdf_loader_missing_file(tmp_path: Path):
    file_path = tmp_path / "missing.pdf"

    loader = PDFLoader()

    with pytest.raises(FileNotFoundError, match="File not found"):
        loader.load(str(file_path))


def test_csv_loader_extracts_rows(tmp_path: Path):
    file_path = tmp_path / "employees.csv"

    file_path.write_text(
        "Name,Department,Role\n"
        "John,IT,Developer\n"
        "Sarah,HR,Manager\n",
        encoding="utf-8",
    )

    loader = CSVLoader()

    result = loader.load(str(file_path))

    assert "Name | Department | Role" in result
    assert "John | IT | Developer" in result
    assert "Sarah | HR | Manager" in result


def test_csv_loader_rejects_non_csv(tmp_path: Path):
    file_path = tmp_path / "employees.txt"
    file_path.write_text("test", encoding="utf-8")

    loader = CSVLoader()

    with pytest.raises(ValueError, match="Expected a CSV file"):
        loader.load(str(file_path))


def test_csv_loader_missing_file(tmp_path: Path):
    file_path = tmp_path / "missing.csv"

    loader = CSVLoader()

    with pytest.raises(FileNotFoundError, match="File not found"):
        loader.load(str(file_path))

def test_document_loader_routes_csv(tmp_path: Path):
    file_path = tmp_path / "employees.csv"

    file_path.write_text(
        "Name,Department,Role\n"
        "John,IT,Developer\n"
        "Sarah,HR,Manager\n",
        encoding="utf-8",
    )

    loader = DocumentLoader()

    result = loader.load(str(file_path))

    assert "Name | Department | Role" in result
    assert "John | IT | Developer" in result
    assert "Sarah | HR | Manager" in result
